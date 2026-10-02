"""Команды админов (ADMIN_IDS): /stats, /users, /stuck. Не-админам не отвечают."""
import html
import shutil
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

from aiogram import Router
from aiogram.filters import Command, Filter
from aiogram.types import BufferedInputFile, Message
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import config
from storage.db import FUNNEL, Database, UserInfo
from storage.files import FileStore

router = Router(name="admin")

# Названия шагов для людей
STEP_TITLES = {
    "start": "/start",
    "instruction_open": "открыл инструкцию",
    "new_analysis": "новый анализ",
    "file_ok": "файл принят",
    "file_error": "ошибка файла",
    "label": "название",
    "price": "цена",
    "drr": "ДРР",
    "result_sent": "получил результат",
    "excel_sent": "получил Excel",
    "excel_download": "скачал Excel из истории",
    "recalc": "пересчёт",
    "history_open": "история",
    "cancel": "отмена",
    "lead_request": "хочет разбор",
    "lead_contact": "оставил контакт",
    "materials_open": "материалы",
}
STUCK_LIST_MAX = 50


class IsAdmin(Filter):
    async def __call__(self, message: Message) -> bool:
        return message.from_user is not None and message.from_user.id in config.ADMIN_IDS


def dir_size(path: Path) -> int:
    if not path.exists():
        return 0
    return sum(p.stat().st_size for p in path.rglob("*") if p.is_file())


def fmt_bytes(n: float) -> str:
    for unit in ("Б", "КБ", "МБ", "ГБ"):
        if n < 1024 or unit == "ГБ":
            return f"{n:.0f} {unit}" if unit == "Б" else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024
    return ""


def _existing(path: Path) -> Path:
    while not path.exists() and path != path.parent:
        path = path.parent
    return path


def step_title(step: str | None) -> str:
    return STEP_TITLES.get(step, step or "—")


def format_funnel(counts: dict[str, int], steps: tuple[str, ...] = FUNNEL) -> list[str]:
    """«start 10 → new_analysis 8 (80 %) → …» — по строке на шаг, % от предыдущего шага."""
    lines, prev = [], None
    for step in steps:
        n = counts.get(step, 0)
        conv = ""
        if prev is not None:
            conv = f" ({round(n / prev * 100)} %)" if prev else " (—)"
        lines.append(f"{step_title(step)}: {n}{conv}")
        prev = n
    return lines


async def funnel_block(db: Database, days: int, now: datetime) -> str:
    since = now - timedelta(days=days)
    counts = await db.funnel(since)
    materials = await db.count_users_with_event("materials_open", since)
    errors = await db.top_file_errors(since)
    lines = [f"<b>Воронка за {days} дн.</b> (уникальные пользователи)"]
    lines += [html.escape(line) for line in format_funnel(counts)]
    lines.append(f"📚 Открыли материалы: {materials}")
    if errors:
        lines.append("Топ ошибок файла:")
        lines += [f"• {html.escape(reason)} — {n}" for reason, n in errors]
    return "\n".join(lines)


@router.message(Command("stats"), IsAdmin())
async def cmd_stats(message: Message, db: Database, files: FileStore) -> None:
    now = datetime.now()
    s = await db.stats()
    db_path = Path(db.path)
    uploads = dir_size(files.root)
    db_size = db_path.stat().st_size if db_path.exists() else 0
    free = shutil.disk_usage(_existing(files.root)).free
    mode = ("открытый" if not config.ALLOWED_USER_IDS
            else f"whitelist ({len(config.ALLOWED_USER_IDS)} ID)")
    await message.answer(
        "<b>📈 Статистика</b>\n\n"
        f"Режим доступа: {mode}\n"
        f"Пользователей: {s.users}\n"
        f"Анализов за сутки: {s.analyses_day}\n"
        f"Анализов за неделю: {s.analyses_week}\n"
        f"Анализов всего: {s.analyses_total}\n\n"
        f"Файлов выгрузок: {s.uploads_with_file} · {fmt_bytes(uploads)}\n"
        f"База: {fmt_bytes(db_size)}\n"
        f"Свободно на диске: {fmt_bytes(free)}\n"
        f"Файлы хранятся {config.FILE_TTL_DAYS} дн.\n\n"
        + await funnel_block(db, 7, now) + "\n\n" + await funnel_block(db, 30, now)
    )


# ---------- /users ----------

USERS_HEADERS = ["ID", "Username", "Имя", "Телефон", "Первый визит", "Последний визит",
                 "Анализов", "Последний шаг", "Когда", "Застрял"]


def _fmt_dt(d: datetime | None) -> str:
    return d.strftime("%d.%m.%Y %H:%M") if d else ""


def build_users_excel(users: list[UserInfo], now: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"
    ws.append(USERS_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    idle = timedelta(minutes=config.STUCK_MINUTES)
    for u in users:
        ws.append([u.user_id, f"@{u.username}" if u.username else "", u.full_name, u.phone or "",
                   _fmt_dt(u.first_seen), _fmt_dt(u.last_seen), u.analyses,
                   step_title(u.last_step) if u.last_step else "", _fmt_dt(u.last_step_at),
                   "да" if u.is_stuck(now, idle) else "нет"])
        if u.username:
            cell = ws.cell(row=ws.max_row, column=2)
            cell.hyperlink = f"https://t.me/{u.username}"
            cell.style = "Hyperlink"
    for col, w in zip("ABCDEFGHIJ", (13, 22, 26, 16, 17, 17, 10, 20, 17, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.message(Command("users"), IsAdmin())
async def cmd_users(message: Message, db: Database) -> None:
    now = datetime.now()
    users = await db.all_users()
    data = build_users_excel(users, now)
    await message.answer_document(
        BufferedInputFile(data, filename=f"users_{now:%Y-%m-%d}.xlsx"),
        caption=f"👥 Пользователей: {len(users)}")


# ---------- /stuck ----------

async def stuck_users(db: Database, now: datetime) -> list[UserInfo]:
    idle = timedelta(minutes=config.STUCK_MINUTES)
    since = now - timedelta(days=config.STUCK_DAYS)
    users = [u for u in await db.all_users()
             if u.is_stuck(now, idle) and u.last_step_at and u.last_step_at >= since]
    return sorted(users, key=lambda u: u.last_step_at, reverse=True)


@router.message(Command("stuck"), IsAdmin())
async def cmd_stuck(message: Message, db: Database) -> None:
    users = await stuck_users(db, datetime.now())
    if not users:
        await message.answer(f"За {config.STUCK_DAYS} дн. никто не застрял 👌")
        return
    lines = [f"<b>🧱 Застряли за {config.STUCK_DAYS} дн.:</b> {len(users)}", ""]
    for u in users[:STUCK_LIST_MAX]:
        lines.append(f"• {html.escape(u.full_name)} {html.escape(u.mention)} — "
                     f"{html.escape(step_title(u.last_step))}, {_fmt_dt(u.last_step_at)}")
    if len(users) > STUCK_LIST_MAX:
        lines.append(f"…и ещё {len(users) - STUCK_LIST_MAX} — в /users")
    await message.answer("\n".join(lines))
