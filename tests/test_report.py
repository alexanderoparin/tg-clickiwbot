from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.analyzer import analyze
from core.report import (MAIN_HEADERS, ReportMeta, build_excel, clicks_word,
                         excel_filename, fmt_num, format_summary, keys_word, plural)


@pytest.fixture(scope="module")
def ctx(parsed):
    res = analyze(parsed.clusters, 9500, 4)
    meta = ReportMeta(analysis_id=7, label="Бомбер <кожаный>", orig_name="file.xlsx")
    return res, meta


@pytest.fixture(scope="module")
def wb(ctx):
    res, meta = ctx
    return load_workbook(BytesIO(build_excel(res, meta)))


# ---------- сводка в чат ----------

def test_summary(ctx):
    res, meta = ctx
    s = format_summary(res, meta)
    assert s.startswith("<b>📊 Анализ: Бомбер &lt;кожаный&gt;</b>\n")
    assert "Цена 9 500 ₽ · норма ДРР 4 % · допустимый CPO 380 ₽" in s
    assert "Расход на поиск: 4 533 ₽ · ДРР поиска 2,1 %" in s
    assert "🔴 Возможно удалить: 1 ключ · 409 ₽ · 94 клика · 5 корз. · 0 зак.\n" in s
    assert "🟡 Работать: 0\n" in s
    assert "🟢 Норма: 13 ключей · 2 258 ₽ · 519 кликов · 58 корз. · 23 зак. · ДРР 1 %" in s
    assert "⚪ Мало данных: 138 ключей · 1 866 ₽ · 429 кликов · 41 корз. · 0 зак. (41 % расхода)" in s
    assert "<b>Возможно удалить:</b>\n1. косуха женская — 409 ₽, 94 клика, 5 корз., 0 зак." in s
    assert " кл." not in s, "«кл.» больше не используется ни для ключей, ни для кликов"
    assert s.endswith("💡 Много кластеров без статистики — загрузи статистику за 30 дней")
    assert "14–30" not in s


def test_summary_removed_parts(ctx):
    res, meta = ctx
    s = format_summary(res, meta)
    for gone in ("Топ на удаление", "🔴 Удалить", "Экономия", "💰", "⚫", "релевант",
                 "период не определён", "17.09", "23.09"):
        assert gone not in s, gone


def test_summary_without_label_and_yellow_drr(parsed):
    s = format_summary(analyze(parsed.clusters, 9500, 3), ReportMeta())
    assert s.startswith("<b>📊 Анализ</b>\n")
    assert "🟡 Работать: 1 ключ · 1 536 ₽ · 353 клика · 27 корз. · 5 зак. · ДРР 3,2 %" in s
    assert "1. кожаная куртка для женщин — 1 536 ₽, 353 клика, 27 корз., 5 зак., ДРР 3,2 %" in s


@pytest.mark.parametrize("n,word", [
    (1, "1 ключ"), (2, "2 ключа"), (4, "4 ключа"), (5, "5 ключей"), (11, "11 ключей"),
    (12, "12 ключей"), (14, "14 ключей"), (21, "21 ключ"), (23, "23 ключа"), (25, "25 ключей"),
    (111, "111 ключей"), (1021, "1 021 ключ"),
])
def test_keys_word(n, word):
    assert keys_word(n) == word


@pytest.mark.parametrize("n,word", [
    (1, "1 клик"), (2, "2 клика"), (5, "5 кликов"), (11, "11 кликов"), (21, "21 клик"),
    (23, "23 клика"), (476, "476 кликов"), (94.0, "94 клика"),
])
def test_clicks_word(n, word):
    assert clicks_word(n) == word


def _section(s: str, title: str) -> list[str]:
    return s.split(f"<b>{title}</b>\n")[1].split("\n\n")[0].splitlines()


def test_summary_five_examples(parsed):
    # цена 100 ₽ и ДРР 1 %: 139 ключей в 🔴 и 13 в 🟡 — примеры обрезаются до 5
    s = format_summary(analyze(parsed.clusters, 100, 1), ReportMeta())
    red = _section(s, "Возможно удалить:")
    yel = _section(s, "Работать (снизить ставку, проверить карточку):")
    assert [line.split(".")[0] for line in red[:5]] == ["1", "2", "3", "4", "5"]
    assert red[5:] == ["…и ещё 134 — в Excel"]
    assert [line.split(".")[0] for line in yel[:5]] == ["1", "2", "3", "4", "5"]
    assert yel[5:] == ["…и ещё 8 — в Excel"]
    assert "🔴 Возможно удалить: 139 ключей ·" in s and "🟡 Работать: 13 ключей ·" in s


def test_summary_no_tail_when_five_or_less(ctx):
    res, meta = ctx
    assert "…и ещё" not in format_summary(res, meta)


# ---------- Excel ----------

def test_excel_sheets_and_no_flag(wb):
    assert wb.sheetnames == ["Анализ", "Минус-фразы", "Параметры"]
    ws = wb["Анализ"]
    assert [c.value for c in ws[5]][:10] == MAIN_HEADERS, "шапка таблицы — строка 5"
    assert "Флаг ⚫" not in [c.value for c in ws[5]]
    # 152 ключа + 2 разделителя между тремя блоками = строки 6–159
    assert ws.auto_filter.ref == "A5:J159", "автофильтр только на таблице, со строки 5"
    assert ws.freeze_panes == "A6", "шапка с логотипом и заголовки остаются наверху"
    params = {r[0]: r[1] for r in wb["Параметры"].iter_rows(min_row=5, values_only=True)}
    assert "Период" not in params and not any("Экономия" in k for k in params)


def test_excel_main_table(wb):
    ws = wb["Анализ"]
    assert ws["A6"].value == "косуха женская"
    assert ws["B6"].value == "🔴 Возможно удалить"
    assert ws["A6"].fill.fgColor.rgb.endswith("FFC7CE")


# ---------- фирменная шапка ----------

def _logo_px(ws) -> tuple[float, float]:
    # размер на листе хранится в якоре (EMU); img.width при чтении — исходные px картинки
    ext = ws._images[0].anchor.ext
    return ext.width / 9525, ext.height / 9525


def test_excel_brand_header(wb, ctx):
    res, meta = ctx
    ws = wb["Анализ"]
    assert len(ws._images) == 1, "логотип на листе"
    anchor = ws._images[0].anchor
    assert _logo_px(ws) == (90, 90), "квадрат 90×90 px"
    assert anchor._from.col == 0 and anchor._from.row == 0, "якорь A1"
    heights = [ws.row_dimensions[r].height for r in (1, 2, 3, 4)]
    assert heights == [18, 18, 18, 18]
    assert sum(heights) * 96 / 72 >= 90, "строки 1–4 вмещают логотип и он не налезает на строку 5"
    b1 = ws["B1"]
    assert b1.value == "Clicki" and b1.hyperlink.target == "https://click-i.ru"
    assert b1.font.bold and b1.font.size == 16
    assert ws["B2"].value == "Анализ рекламных ключей WB по ДРР" and ws["B2"].font.size == 11
    assert ws["B2"].font.color.rgb.endswith("808080")
    assert ws["B3"].value == (f"Бомбер <кожаный> · цена 9 500 ₽ · норма ДРР 4 % · "
                              f"{meta.created_at:%d.%m.%Y}")
    assert ws["B3"].font.size == 10 and ws["B3"].font.color.rgb.endswith("808080")
    # справа от шапки и в строке 4 пусто — ничего не налезает на логотип и текст
    assert all(ws.cell(row=r, column=c).value is None for r in (1, 2, 3) for c in range(3, 21))
    assert all(ws.cell(row=4, column=c).value is None for c in range(1, 21))


def test_excel_no_empty_row_between_header_and_data(wb):
    # пустой строки 5 больше нет: сразу после шапки (1–4) идут данные
    ws = wb["Анализ"]
    assert ws["A5"].value == "Кластер" and ws["L5"].value == "Блок"
    assert wb["Минус-фразы"]["A5"].value == "Минус-фразы"
    assert wb["Параметры"]["A5"].value == "Цена товара, ₽"


def test_excel_other_sheets_brand_and_data_from_row_5(wb):
    for name in ("Минус-фразы", "Параметры"):
        ws = wb[name]
        assert len(ws._images) == 1 and _logo_px(ws) == (90, 90), name
        assert [ws.row_dimensions[r].height for r in (1, 2, 3, 4)] == [18, 18, 18, 18], name
        assert ws["B1"].value == "Clicki" and ws["B1"].hyperlink.target == "https://click-i.ru"
        assert all(ws.cell(row=r, column=1).value is None for r in (1, 2, 3, 4)), name
    minus = wb["Минус-фразы"]
    assert minus["A5"].value == "Минус-фразы"
    assert [c.value for c in minus["A"]][5:] == ["косуха женская"]
    assert wb["Параметры"]["A5"].value == "Цена товара, ₽" and wb["Параметры"]["B5"].value == 9500


def test_minus_phrases_single_solid_column(parsed):
    # много фраз (цена 100 ₽, ДРР 1 %: 139 🔴) — один сплошной столбец без пустых строк
    ws = load_workbook(BytesIO(build_excel(analyze(parsed.clusters, 100, 1), ReportMeta())))["Минус-фразы"]
    col = [ws.cell(row=r, column=1).value for r in range(6, ws.max_row + 1)]
    assert len(col) == 139 and all(col)
    assert all(ws.cell(row=r, column=c).value is None
               for r in range(5, ws.max_row + 1) for c in range(2, 5)), "рядом ничего нет"


def test_excel_without_logo(ctx, tmp_path, caplog):
    res, meta = ctx
    with caplog.at_level("WARNING"):
        data = build_excel(res, meta, logo_path=tmp_path / "нет_такого.png")
    wb = load_workbook(BytesIO(data))
    for ws in wb:
        assert ws._images == [] and ws["B1"].value == "Clicki"
    assert wb["Анализ"]["A6"].value == "косуха женская", "отчёт собрался как обычно"
    assert "Нет логотипа" in caplog.text


def test_real_logo_asset():
    import config
    from PIL import Image
    im = Image.open(config.LOGO_PATH)
    assert im.width == im.height, "логотип квадратный"
    assert im.size == (260, 260) and im.mode == "RGBA"
    assert im.getpixel((0, 0))[3] == 0, "скруглённые углы прозрачные"
    assert im.getpixel((130, 130))[3] == 255


def test_logo_square_90_on_all_sheets(wb):
    for ws in wb:
        assert len(ws._images) == 1, ws.title
        anchor = ws._images[0].anchor
        assert _logo_px(ws) == (90, 90), f"{ws.title}: квадрат 90×90, пропорции 1:1"
        assert (anchor._from.col, anchor._from.row) == (0, 0), f"{ws.title}: якорь A1"


HDR = ["Блок", "Ключей", "Затраты, ₽", "Доля расхода", "Клики", "Корзины", "Заказы", "CPO", "ДРР %"]


def _side(ws, r):
    return [ws.cell(row=r, column=c).value for c in range(12, 21)]


def _main(ws, r):
    return [ws.cell(row=r, column=c).value for c in range(1, 11)]


def _is_header(ws, r):
    cells = [ws.cell(row=r, column=c) for c in range(12, 21)]
    return ([c.value for c in cells] == HDR and all(c.font.bold for c in cells)
            and all(c.fill.fgColor.rgb.endswith("D9D9D9") for c in cells))


def test_excel_block_summary(wb):
    ws = wb["Анализ"]
    assert ws["K5"].value is None, "одна пустая колонка после «Ср. позиция»"
    # 5 шапка | 6 🔴 (1 ключ) | 7 разделитель | 8–145 ⚪ (138) | 146 разделитель |
    # 147–159 🟢 (13) | 160 заголовок «Итого» | 161 «Итого»
    for r in (5, 7, 146, 160):
        assert _is_header(ws, r), f"строка {r} — заголовок итогов"
    assert "Кластеров" not in _side(ws, 5) and "% расхода" not in _side(ws, 5)

    assert _side(ws, 6) == ["🔴 Возможно удалить", 1, 408.9, 9.0, 94, 5, 0, "—", "—"]
    assert _side(ws, 8) == ["⚪ Мало данных", 138, 1866.15, 41.2, 429, 41, 0, "—", "—"]
    assert _side(ws, 147) == ["🟢 Норма", 13, 2257.65, 49.8, 519, 58, 23, 98.16, 1.03]
    assert _side(ws, 161) == ["Итого", 152, 4532.7, 100.0, 1042, 104, 23, 197.07, 2.07]
    # итог — в первой строке своего блока
    assert ws["B6"].value == "🔴 Возможно удалить"
    assert ws["B8"].value == "⚪ Мало данных" and ws["B145"].value == "⚪ Мало данных"
    assert ws["B147"].value == "🟢 Норма" and ws["B159"].value == "🟢 Норма"
    # заливка итогов — цвет блока, «Итого» — жирным на сером
    assert ws.cell(row=6, column=12).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(row=8, column=20).fill.fgColor.rgb.endswith("EDEDED")
    assert ws.cell(row=147, column=15).fill.fgColor.rgb.endswith("C6EFCE")
    assert ws.cell(row=161, column=12).font.bold
    # кроме заголовков и итогов справа пусто; 🟡 (0 ключей) не выводится
    filled = {r for r in range(1, 165) if any(v is not None for v in _side(ws, r))}
    assert filled == {5, 6, 7, 8, 146, 147, 160, 161}


def test_excel_separator_rows(wb):
    ws = wb["Анализ"]
    for r in (7, 146):
        assert all(v is None for v in _main(ws, r)), f"строка {r} — пустой разделитель"
    # данные не потерялись: 152 ключа по порядку 🔴 → ⚪ → 🟢
    statuses = [ws.cell(row=r, column=2).value for r in range(6, 160)]
    keys = [s for s in statuses if s]
    assert len(keys) == 152 and statuses.count(None) == 2
    assert keys == ["🔴 Возможно удалить"] + ["⚪ Мало данных"] * 138 + ["🟢 Норма"] * 13
    # заливка строк — по новым номерам: ключи залиты, разделители нет
    assert ws["A8"].fill.fgColor.rgb.endswith("EDEDED") and ws["J159"].fill.fgColor.rgb.endswith("C6EFCE")
    assert ws["A7"].fill.fill_type is None


def test_excel_last_row(ctx):
    # на свежей книге: ws.cell() в других тестах создаёт пустые ячейки и сдвигает max_row
    res, meta = ctx
    ws = load_workbook(BytesIO(build_excel(res, meta)))["Анализ"]
    assert ws.max_row == 161, "последняя строка листа — «Итого»"


def test_excel_single_key_blocks(parsed):
    # ДРР 3 %: 🔴 и 🟡 — по одному ключу подряд; заголовки не должны налезать на итоги
    wb = load_workbook(BytesIO(build_excel(analyze(parsed.clusters, 9500, 3), ReportMeta())))
    ws = wb["Анализ"]
    # 5 шапка | 6 🔴 | 7 разд. | 8 🟡 | 9 разд. | 10–147 ⚪ | 148 разд. | 149–160 🟢 | 161–162 Итого
    for r in (5, 7, 9, 148, 161):
        assert _is_header(ws, r), f"строка {r} — заголовок итогов"
    assert _side(ws, 6)[:2] == ["🔴 Возможно удалить", 1] and ws["B6"].value == "🔴 Возможно удалить"
    assert _side(ws, 8) == ["🟡 Работать", 1, 1535.55, 33.9, 353, 27, 5, 307.11, 3.23]
    assert ws["A8"].value == "кожаная куртка для женщин" and ws["B8"].value == "🟡 Работать"
    assert ws.cell(row=8, column=12).fill.fgColor.rgb.endswith("FFEB9C")
    assert _side(ws, 10)[0] == "⚪ Мало данных" and _side(ws, 149)[0] == "🟢 Норма"
    assert _side(ws, 162)[0] == "Итого"
    assert all(v is None for v in _main(ws, 7)) and all(v is None for v in _main(ws, 9))


def test_excel_one_block_only(parsed):
    # все ключи с расходом в одном блоке: без разделителей, «Итого» с заголовком под таблицей
    d = parsed.clusters[parsed.clusters["Заказанные товары"] > 0]
    ws = load_workbook(BytesIO(build_excel(analyze(d, 9500, 100), ReportMeta())))["Анализ"]
    assert _is_header(ws, 5) and _side(ws, 6)[:2] == ["🟢 Норма", 13]
    assert all(ws.cell(row=r, column=1).value for r in range(6, 19))
    assert _is_header(ws, 19) and _side(ws, 20)[0] == "Итого"
    assert ws.auto_filter.ref == "A5:J18" and ws.freeze_panes == "A6"


# ---------- имя файла ----------

@pytest.mark.parametrize("label,aid,name", [
    ("Бомбер", 7, "Бомбер.xlsx"),
    ('ИП Иванов / арт. 12345678: "кожа"?', 1, "ИП Иванов _ арт. 12345678_ _кожа__.xlsx"),
    ('a\\b/c:d*e?f"g<h>i|j', 1, "a_b_c_d_e_f_g_h_i_j.xlsx"),
    ("x" * 80, 1, "x" * 60 + ".xlsx"),
    ("  Бомбер  ", 1, "Бомбер.xlsx"),
    (None, 12, "Анализ_12.xlsx"),
    ("", 3, "Анализ_3.xlsx"),
    ("   ", 4, "Анализ_4.xlsx"),
])
def test_excel_filename(label, aid, name):
    assert excel_filename(label, aid) == name


def test_formatters():
    assert fmt_num(4) == "4" and fmt_num(2.08) == "2,1" and fmt_num(9500.5, 2) == "9 500,5"
    assert [plural(n, "клик", "клика", "кликов") for n in (1, 3, 5, 11, 21, 94)] == [
        "клик", "клика", "кликов", "кликов", "клик", "клика"]
