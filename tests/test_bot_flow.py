"""Сквозной прогон сценариев бота через Dispatcher с фейковой сессией Telegram."""
from datetime import datetime
from itertools import count
from pathlib import Path

import pytest
from aiogram import Bot
from aiogram.client.session.base import BaseSession
from aiogram.exceptions import TelegramBadRequest
from aiogram.methods import (AnswerCallbackQuery, DeleteMessage, EditMessageReplyMarkup,
                             EditMessageText, GetFile,
                             SendAnimation, SendDocument, SendMessage, TelegramMethod)
from aiogram.types import (Animation, CallbackQuery, Chat, Document, File, FSInputFile, Message,
                           ReplyKeyboardMarkup, Update, User)

import config
from bot import texts as T
from bot.keyboards import BTN_CANCEL, BTN_HELP, BTN_HISTORY, BTN_NEW, BTN_SETTINGS
from bot.materials import Materials
from bot.media import MediaSender
from main import build_dispatcher
from storage.db import Database
from storage.files import FileStore

UID = 42
CHAT = Chat(id=UID, type="private")
USER = User(id=UID, is_bot=False, first_name="T")


class FakeSession(BaseSession):
    def __init__(self, file_bytes: bytes):
        super().__init__()
        self.file_bytes = file_bytes
        self.calls: list[TelegramMethod] = []
        self.ids = count(1000)
        self.fail_animation = False     # send_animation падает всегда
        self.bad_file_ids: set[str] = set()  # эти file_id Telegram «не знает»
        self.uploads: dict[str, int] = {}  # имя файла без расширения → число загрузок
        self.fail_delete = False        # delete_message падает («message can't be deleted»)
        self.sent_ids: dict[int, int] = {}  # id(метода) → message_id отправленного сообщения

    def message_id(self, method) -> int:
        return self.sent_ids[id(method)]

    async def make_request(self, bot, method, timeout=None):
        self.calls.append(method)
        if isinstance(method, DeleteMessage):
            if self.fail_delete:
                raise TelegramBadRequest(method=method,
                                         message="Bad Request: message can't be deleted")
            return True
        result = await self._respond(method)
        if isinstance(result, Message):
            self.sent_ids[id(method)] = result.message_id
        return result

    async def _respond(self, method):
        if isinstance(method, SendAnimation):
            if self.fail_animation or method.animation in self.bad_file_ids:
                raise TelegramBadRequest(method=method, message="Bad Request: wrong file")
            if isinstance(method.animation, str):
                file_id = method.animation
            else:
                stem = Path(method.animation.path).stem
                self.uploads[stem] = self.uploads.get(stem, 0) + 1
                file_id = f"{stem}#{self.uploads[stem]}"
            return Message(message_id=next(self.ids), date=datetime.now(), chat=CHAT,
                           caption=method.caption,
                           animation=Animation(file_id=file_id, file_unique_id="au", width=1,
                                               height=1, duration=1))
        if isinstance(method, SendDocument):
            doc = method.document
            if isinstance(doc, str):
                if doc in self.bad_file_ids:
                    raise TelegramBadRequest(method=method, message="Bad Request: wrong file")
                file_id = doc
            elif isinstance(doc, FSInputFile):  # файлы с диска (материалы) — кешируются
                stem = Path(doc.path).stem
                self.uploads[stem] = self.uploads.get(stem, 0) + 1
                file_id = f"{stem}#{self.uploads[stem]}"
            else:  # Excel-отчёт из памяти
                file_id = "report"
            return Message(message_id=next(self.ids), date=datetime.now(), chat=CHAT,
                           caption=method.caption,
                           document=Document(file_id=file_id, file_unique_id="du"))
        if isinstance(method, SendMessage):
            return Message(message_id=next(self.ids), date=datetime.now(), chat=CHAT,
                           text=method.text)
        if isinstance(method, GetFile):
            return File(file_id=method.file_id, file_unique_id="u", file_path="doc.xlsx")
        if isinstance(method, (EditMessageText, EditMessageReplyMarkup)):
            return True
        return True

    async def stream_content(self, url, headers=None, timeout=30, chunk_size=65536, raise_for_status=True):
        yield self.file_bytes

    async def close(self):
        pass


class Harness:
    def __init__(self, tmp_path, fixture_path):
        self.session = FakeSession(fixture_path.read_bytes())
        self.bot = Bot("42:TEST", session=self.session)
        self.db = Database(tmp_path / "bot.db")
        self.files = FileStore(tmp_path / "uploads")
        self.dp = None
        self.upd = count(1)
        self.fixture_name = fixture_path.name
        # видео-подсказка и кеш file_id — во временной папке теста
        self.video = tmp_path / "assets" / "how_to_download.mp4"
        self.video.parent.mkdir()
        self.video.write_bytes(b"fake mp4 v1")
        self.price_video = tmp_path / "assets" / "how_to_price.mp4"
        self.price_video.write_bytes(b"fake price mp4 v1")
        cache = tmp_path / "data" / "media_cache.json"
        self.media = MediaSender(self.video, cache)
        self.price_media = MediaSender(self.price_video, cache)
        # «📚 Полезные материалы» — два PDF во временной папке, управляемые часы
        self.materials_dir = tmp_path / "assets" / "materials"
        self.materials_dir.mkdir()
        for m in T.MATERIALS:
            (self.materials_dir / m["file"]).write_bytes(b"%PDF-1.4 fake " + m["file"].encode())
        self.clock = [1000.0]
        self.materials = Materials(self.materials_dir, cache, repeat_sec=60,
                                   clock=lambda: self.clock[0])

    async def start(self):
        await self.db.init()
        # роутеры — модульные синглтоны; между тестами отцепляем их от прошлого Dispatcher
        from bot.handlers import admin, fallback, history, new_analysis, settings, start
        from bot import lead
        for m in (start, lead, admin, new_analysis, history, settings, fallback):
            m.router._parent_router = None
        self.dp = build_dispatcher(self.db, self.files, flood_interval=0, media=self.media,
                                  price_media=self.price_media, materials=self.materials)

    async def _feed(self, **kw):
        n = len(self.session.calls)
        await self.dp.feed_update(self.bot, Update(update_id=next(self.upd), **kw))
        return self.session.calls[n:]

    async def text(self, text, user=USER):
        return await self._feed(message=Message(
            message_id=next(self.upd), date=datetime.now(), chat=CHAT, from_user=user, text=text))

    async def doc(self, name=None):
        return await self._feed(message=Message(
            message_id=next(self.upd), date=datetime.now(), chat=CHAT, from_user=USER,
            document=Document(file_id="F1", file_unique_id="U1", file_name=name or self.fixture_name,
                              file_size=1000)))

    async def cb(self, data, user=USER):
        msg = Message(message_id=1, date=datetime.now(), chat=CHAT, from_user=USER, text="x")
        return await self._feed(callback_query=CallbackQuery(
            id=str(next(self.upd)), from_user=user, chat_instance="ci", message=msg, data=data))


def texts(calls, chat_id: int = UID):
    """Тексты сообщений и подписи к анимациям в чат пользователя — в порядке отправки.
    Уведомления админам (в их чаты) сюда не попадают — см. admin_texts()."""
    return [c.text if isinstance(c, SendMessage) else c.caption
            for c in calls if isinstance(c, (SendMessage, SendAnimation))
            and c.chat_id == chat_id]


def admin_texts(calls, admin_id: int) -> list[str]:
    return [c.text for c in calls if isinstance(c, SendMessage) and c.chat_id == admin_id]


def summary(calls) -> str:
    """Сводка анализа из ответа (после неё может прийти запрос контакта)."""
    [s] = [t for t in texts(calls) if t and t.startswith("<b>📊 Анализ")]
    return s


def animations(calls) -> list[SendAnimation]:
    return [c for c in calls if isinstance(c, SendAnimation)]


def buttons(calls) -> dict[str, str]:
    """Текст кнопки → callback_data по всем inline-клавиатурам в ответах."""
    out = {}
    for c in calls:
        kb = getattr(c, "reply_markup", None)
        for row in getattr(kb, "inline_keyboard", None) or []:
            for b in row:
                out[b.text] = b.callback_data
    return out


@pytest.fixture
async def h(tmp_path, fixture_path, monkeypatch):
    monkeypatch.setattr(config, "ALLOWED_USER_IDS", {UID})
    import main
    monkeypatch.setattr(main.config, "ALLOWED_USER_IDS", {UID})
    harness = Harness(tmp_path, fixture_path)
    await harness.start()
    return harness


async def test_access_denied(h):
    stranger = User(id=7, is_bot=False, first_name="S")
    assert texts(await h.text("/start", user=stranger)) == ["Нет доступа"]
    assert "Привет" in texts(await h.text("/start"))[0]


async def new_analysis(h, label="ИП Иванов", price="9 500", drr_btn="4 %"):
    assert "Топ поисковых кластеров" in texts(await h.text(BTN_NEW))[0]
    r = await h.doc()
    assert "334 кластеров" in "\n".join(texts(r))
    if label:
        r = await h.text(label)
    else:
        r = await h.cb("skip_label")
    assert "цену реализации" in texts(r)[-1]
    r = await h.text(price)
    assert "ДРР" in texts(r)[-1]
    return await h.cb(buttons(r)[drr_btn])


async def test_full_flow(h):
    r = await new_analysis(h)
    s = summary(r)
    assert s.startswith("<b>📊 Анализ: ИП Иванов</b>\n")
    assert "🔴 Возможно удалить: 1 ключ · 409 ₽ · 94 клика" in s
    assert "⚫" not in s and "Экономия" not in s
    docs = [c for c in r if isinstance(c, SendDocument)]
    assert docs and docs[0].document.filename == "ИП Иванов.xlsx"
    assert await h.db.count_analyses(UID) == 1

    # Excel приходит сам сразу после сводки — кнопки «📥 Excel» под сводкой нет
    assert "📥 Excel" not in buttons(r)
    kinds = [type(c) for c in r if isinstance(c, (SendMessage, SendDocument))]
    i_sum = texts(r).index(s)
    assert kinds[i_sum:i_sum + 2] == [SendMessage, SendDocument], "Excel сразу за сводкой"

    # пересчёт: файл не просим, создаётся новая запись на тот же upload
    r = await h.cb("an:recalc:1")
    assert "Файл тот же" in texts(r)[0]
    r = await h.text("9500,50")
    r = await h.cb(buttons(r)["Свой"])
    assert "свой" in texts(r)[0].lower()
    assert "ДРР должна быть" in texts(await h.text("500"))[0]
    r = await h.text("3")
    assert "🟡 Работать: 1" in summary(r)
    items = await h.db.list_analyses(UID)
    assert [a.id for a in items] == [2, 1] and items[0].upload_id == items[1].upload_id
    assert items[0].label == "ИП Иванов"

    # история
    r = await h.text(BTN_HISTORY)
    hist = texts(r)[0]
    assert "ИП Иванов · 9 500,5 ₽ · ДРР 3 % · 🔴1 🟡1" in hist
    assert "17.09" not in hist
    r = await h.cb(buttons(r)["2"])  # второй в списке = анализ №1
    assert "норма ДРР 4 %" in texts(r)[0]
    b = buttons(r)
    r = await h.cb(b["🗑 Удалить"])
    r = await h.cb(buttons(r)["Да, удалить"])
    assert await h.db.count_analyses(UID) == 1
    upload_path = (await h.db.list_analyses(UID))[0].file_path
    from pathlib import Path
    assert Path(upload_path).exists()  # файл ещё нужен анализу №2
    await h.cb("an:del:2")
    await h.cb("an:delok:2")
    assert not Path(upload_path).exists()
    assert "История пуста" in texts(await h.text(BTN_HISTORY))[0]


async def test_cancel_and_bad_input(h):
    await h.text(BTN_NEW)
    assert "не Excel" in texts(await h.doc("report.csv"))[0]
    assert "Топ поисковых" in texts(await h.text("привет"))[0]
    assert texts(await h.text(BTN_CANCEL)) == ["Отменено."]
    # после отмены текст не воспринимается как шаг FSM
    assert "меню" in texts(await h.text("9500"))[0]

    await h.text(BTN_NEW)
    await h.doc()
    await h.cb("skip_label")
    assert "Не понял цену" in texts(await h.text("дорого"))[0]
    assert "от 1 до 10 000 000" in texts(await h.text("0"))[0]
    # кнопка меню на шаге FSM срабатывает как меню, а не как ввод
    assert "История" in texts(await h.text(BTN_HISTORY))[0]
    assert await h.db.count_analyses(UID) == 0


async def test_settings_defaults(h):
    r = await h.text(BTN_SETTINGS)
    assert "не задана" in texts(r)[0]
    await h.cb(buttons(r)["Изменить цену"])
    await h.text("9 500")
    r = await h.cb("set:drr")
    await h.cb(buttons(r)["5 %"])
    s = await h.db.get_settings(UID)
    assert (s.default_price, s.default_drr) == (9500, 5)

    await h.text(BTN_NEW)
    await h.doc()
    r = await h.cb("skip_label")
    b = buttons(r)
    assert "Взять 9 500 ₽" in b
    r = await h.cb(b["Взять 9 500 ₽"])
    assert "⭐ 5 %" in buttons(r)
    r = await h.cb(buttons(r)["⭐ 5 %"])
    assert "норма ДРР 5 %" in summary(r)
    assert summary(r).startswith("<b>📊 Анализ</b>\n"), "название пропущено"
    docs = [c for c in r if isinstance(c, SendDocument)]
    assert docs[0].document.filename == "Анализ_1.xlsx"


async def test_stale_button(h):
    r = await h.cb("drr:4")
    ans = [c for c in r if isinstance(c, AnswerCallbackQuery)]
    assert ans and "устарела" in ans[0].text


# ---------- доступ и ограничения ----------

STRANGER = User(id=7, is_bot=False, first_name="S")
ADMIN = User(id=1, is_bot=False, first_name="A")


async def rebuild(h, monkeypatch, allowed=None, admins=None, flood=0):
    """Пересобрать Dispatcher с другими настройками доступа."""
    if allowed is not None:
        monkeypatch.setattr(config, "ALLOWED_USER_IDS", allowed)
    if admins is not None:
        monkeypatch.setattr(config, "ADMIN_IDS", admins)
    from bot.handlers import admin, fallback, history, new_analysis, settings, start
    from bot import lead
    for m in (start, lead, admin, new_analysis, history, settings, fallback):
        m.router._parent_router = None
    h.dp = build_dispatcher(h.db, h.files, flood_interval=flood, media=h.media,
                            price_media=h.price_media, materials=h.materials)


async def test_open_mode(h, monkeypatch):
    await rebuild(h, monkeypatch, allowed=set())
    assert "Привет" in texts(await h.text("/start", user=STRANGER))[0]


async def test_admin_stats(h, monkeypatch):
    await rebuild(h, monkeypatch, allowed={UID}, admins={1})
    await new_analysis(h)
    stats = texts(await h.text("/stats", user=ADMIN))[0]
    assert "Режим доступа: whitelist (1 ID)" in stats
    # пользователь + админ, который отправил /stats: оба теперь в базе пользователей
    assert "Пользователей: 2" in stats and "Анализов за сутки: 1" in stats
    assert "Анализов за неделю: 1" in stats and "Свободно на диске" in stats
    # не админ: команда не раскрывается
    assert texts(await h.text("/stats")) == ["Выбери действие в меню."]


async def test_daily_limit(h, monkeypatch):
    monkeypatch.setattr(config, "MAX_ANALYSES_PER_DAY", 1)
    await new_analysis(h)
    assert "Лимит: 1 анализов за сутки" in texts(await h.text(BTN_NEW))[0]
    assert "Лимит" in texts(await h.cb("an:recalc:1"))[0]
    assert await h.db.count_analyses(UID) == 1


async def test_file_too_big(h):
    await h.text(BTN_NEW)
    big = Message(message_id=999, date=datetime.now(), chat=CHAT, from_user=USER,
                  document=Document(file_id="F", file_unique_id="U", file_name=h.fixture_name,
                                    file_size=10 * 1024 * 1024 + 1))
    r = await h._feed(message=big)
    assert "больше 10 МБ" in texts(r)[0]


async def test_expired_file(h):
    from datetime import timedelta
    from pathlib import Path
    await new_analysis(h)
    path = Path((await h.db.list_analyses(UID))[0].file_path)
    service = h.dp["service"]
    assert await service.cleanup_old_files(now=datetime.now() + timedelta(days=59)) == 0
    assert path.exists()
    assert await service.cleanup_old_files(now=datetime.now() + timedelta(days=61)) == 1
    assert not path.exists()

    assert texts(await h.cb("an:xl:1")) == ["Файл устарел, загрузи заново."]
    assert texts(await h.cb("an:recalc:1")) == ["Файл устарел, загрузи заново."]
    opened = texts(await h.cb("an:open:1"))[0]
    assert "⌛ Исходный файл удалён" in opened and "🔴 Возможно удалить: 1 · 409 ₽" in opened
    assert "ИП Иванов" in texts(await h.text(BTN_HISTORY))[0], "запись в истории осталась"


async def test_flood(h, monkeypatch):
    await rebuild(h, monkeypatch, flood=1.0)
    assert "Привет" in texts(await h.text("/start"))[0]
    assert "Не так быстро" in texts(await h.text("/start"))[0]
    assert texts(await h.text("/start")) == []


# ---------- видео-подсказки ----------

DL = "how_to_download"
PR = "how_to_price"


def _cache(h) -> dict:
    import json
    return json.loads(h.media.cache_path.read_text(encoding="utf-8"))


def _is_upload(a: SendAnimation) -> bool:
    return isinstance(a.animation, FSInputFile)


async def test_video_on_new_analysis_then_by_file_id(h):
    r = await h.text(BTN_NEW)
    [a] = animations(r)
    assert _is_upload(a), "первый раз — загрузка файла"
    assert a.caption == T.ASK_FILE
    assert "📅 Выгружай статистику за 30 дней" in a.caption and "Как скачать — на видео 👆" in a.caption
    assert isinstance(a.reply_markup, ReplyKeyboardMarkup)
    assert a.reply_markup.keyboard[0][0].text == BTN_CANCEL, "«❌ Отмена» под видео"
    assert not [c for c in r if isinstance(c, SendMessage)], "текст — только в подписи"
    assert _cache(h)[DL]["file_id"] == f"{DL}#1"

    assert texts(await h.text(BTN_CANCEL)) == ["Отменено."]

    [a2] = animations(await h.text(BTN_NEW))
    assert a2.animation == f"{DL}#1" and h.session.uploads[DL] == 1
    assert "334 кластеров" in "\n".join(texts(await h.doc()))


async def test_price_step_video_then_by_file_id(h):
    await h.text(BTN_NEW)
    await h.doc()
    r = await h.text("Бомбер")
    [a] = animations(r)
    assert _is_upload(a) and a.caption == T.ASK_PRICE
    assert "<b>цену реализации</b>" in a.caption and "<b>«Цена со скидкой»</b> 👆" in a.caption
    assert "«Товары и цены» → «Цены и скидки»" in a.caption
    assert a.reply_markup is None, "цены по умолчанию нет — кнопки «Взять» нет"
    assert _cache(h)[PR]["file_id"] == f"{PR}#1"
    # «❌ Отмена» на шаге цены работает как раньше
    assert texts(await h.text(BTN_CANCEL)) == ["Отменено."]

    # второй раз — по file_id
    await h.text(BTN_NEW)
    await h.doc()
    [a2] = animations(await h.cb("skip_label"))
    assert a2.animation == f"{PR}#1" and h.session.uploads[PR] == 1
    assert "ДРР" in texts(await h.text("9 500"))[-1], "цена принимается, дальше шаг ДРР"


async def test_price_step_default_price_button(h):
    await h.db.set_default(UID, price=9500)
    await h.text(BTN_NEW)
    await h.doc()
    [a] = animations(await h.cb("skip_label"))
    assert [[b.text for b in row] for row in a.reply_markup.inline_keyboard] == [["Взять 9 500 ₽"]]
    r = await h.cb(buttons([a])["Взять 9 500 ₽"])
    assert "ДРР" in texts(r)[-1]


async def test_recalc_shows_price_video(h):
    await new_analysis(h)
    r = await h.cb("an:recalc:1")
    assert "Файл тот же" in texts(r)[0]
    [a] = animations(r)
    assert a.caption == T.ASK_PRICE and a.animation == f"{PR}#1"


async def test_video_in_instruction_and_help(h):
    for trigger in (BTN_HELP, "/help"):
        r = await h.text(trigger)
        sent = [c for c in r if isinstance(c, (SendAnimation, SendMessage))]
        assert [type(c) for c in sent] == [SendAnimation, SendAnimation, SendMessage]
        assert sent[0].caption == "Как скачать выгрузку из кабинета WB"
        assert sent[1].caption == T.PRICE_CAPTION
        assert sent[2].text == T.INSTRUCTION
    anims = [c for c in h.session.calls if isinstance(c, SendAnimation)]
    assert [_is_upload(a) for a in anims] == [True, True, False, False]
    assert [a.animation for a in anims[2:]] == [f"{DL}#1", f"{PR}#1"]
    assert h.session.uploads == {DL: 1, PR: 1}


def test_instruction_text():
    ins = T.INSTRUCTION
    assert "Продвижение → Кампании → открой нужную кампанию" in ins
    assert "выбери период за 30 дней → «Скачать»" in ins
    assert "Как это выглядит — на видео выше." in ins
    assert ("<b>4. Укажи цену реализации</b>\n"
            "Это «Цена со скидкой» в разделе «Товары и цены» → «Цены и скидки». С неё WB считает "
            "ДРР, поэтому цифры бота совпадут с кабинетом.\n\n<b>5.") in ins
    assert "Не бери" not in ins and "Не путай" not in ins
    assert "{" not in ins and "14–30" not in ins
    assert len(ins) <= 4096
    for cap in (T.ASK_FILE, T.ASK_PRICE, T.HOWTO_CAPTION, T.PRICE_CAPTION):
        assert len(cap) <= 1024


def test_no_real_client_names_in_texts():
    import inspect
    import bot.handlers.new_analysis as na
    src = inspect.getsource(T) + inspect.getsource(na)
    assert "ИП Иванов / арт. 12345678" in T.ASK_LABEL and na.ASK_LABEL == T.ASK_LABEL
    for real in ("Морозов", "146910"):
        assert real not in src


async def test_label_step_example(h):
    await h.text(BTN_NEW)
    r = await h.doc()
    assert "Например, <i>ИП Иванов / арт. 12345678</i>" in texts(r)[-1]


async def test_video_cache_shared_between_places(h):
    await h.text(BTN_HELP)
    [a] = animations(await h.text(BTN_NEW))
    assert a.animation == f"{DL}#1", "кеш общий для инструкции и шага «пришли файл»"
    await h.doc()
    [a] = animations(await h.cb("skip_label"))
    assert a.animation == f"{PR}#1", "и для инструкции и шага цены"


async def test_video_missing_sends_text(h, caplog):
    h.video.unlink()
    h.price_video.unlink()
    with caplog.at_level("WARNING"):
        r = await h.text(BTN_NEW)
    assert not animations(r)
    assert texts(r) == [T.ASK_FILE]
    msg = [c for c in r if isinstance(c, SendMessage)][0]
    assert msg.reply_markup.keyboard[0][0].text == BTN_CANCEL
    assert "Нет видео-подсказки" in caplog.text
    await h.doc()
    r = await h.cb("skip_label")
    assert not animations(r) and texts(r)[-1] == T.ASK_PRICE, "шаг цены — текстом"
    assert "ДРР" in texts(await h.text("9 500"))[-1]
    r = await h.text(BTN_HELP)
    assert texts(r) == [T.HOWTO_CAPTION, T.PRICE_CAPTION, T.INSTRUCTION]


async def test_one_video_missing_in_instruction(h):
    h.price_video.unlink()
    r = await h.text(BTN_HELP)
    sent = [c for c in r if isinstance(c, (SendAnimation, SendMessage))]
    assert [type(c) for c in sent] == [SendAnimation, SendMessage, SendMessage]
    assert [texts([c])[0] for c in sent] == [T.HOWTO_CAPTION, T.PRICE_CAPTION, T.INSTRUCTION]


async def test_video_send_error_sends_text(h, caplog):
    h.session.fail_animation = True
    with caplog.at_level("WARNING"):
        r = await h.text(BTN_NEW)
    assert [c.caption for c in animations(r)] == [T.ASK_FILE], "попытка была"
    assert [c.text for c in r if isinstance(c, SendMessage)] == [T.ASK_FILE], "и текст вместо неё"
    assert "Не удалось отправить видео-подсказку" in caplog.text
    assert not h.media.cache_path.exists() or DL not in _cache(h)
    assert "334 кластеров" in "\n".join(texts(await h.doc()))
    r = await h.cb("skip_label")
    assert [c.text for c in r if isinstance(c, SendMessage)] == [T.ASK_PRICE]


async def test_video_changed_file_resets_cache(h):
    await h.text(BTN_HELP)
    assert _cache(h)[DL]["file_id"] == f"{DL}#1"
    h.video.write_bytes(b"fake mp4 v2 - other size")
    anims = animations(await h.text(BTN_HELP))
    assert _is_upload(anims[0]), "файл изменился — загрузка заново"
    assert anims[1].animation == f"{PR}#1", "второе видео не менялось — по file_id"
    assert _cache(h)[DL]["file_id"] == f"{DL}#2"
    assert _cache(h)[DL]["size"] == h.video.stat().st_size
    assert animations(await h.text(BTN_HELP))[0].animation == f"{DL}#2"


async def test_video_changed_mtime_resets_cache(h):
    import os
    await h.text(BTN_HELP)
    st = h.price_video.stat()
    os.utime(h.price_video, ns=(st.st_atime_ns, st.st_mtime_ns + 10**9))
    anims = animations(await h.text(BTN_HELP))
    assert not _is_upload(anims[0]) and _is_upload(anims[1])


async def test_video_stale_file_id_reuploads(h):
    await h.text(BTN_HELP)
    h.session.bad_file_ids.add(f"{DL}#1")
    r = await h.text(BTN_HELP)
    anims = [a for a in animations(r) if a.caption == T.HOWTO_CAPTION]
    assert [type(a.animation) for a in anims] == [str, FSInputFile]
    assert _cache(h)[DL]["file_id"] == f"{DL}#2"
    assert T.HOWTO_CAPTION not in [c.text for c in r if isinstance(c, SendMessage)]


def test_real_video_assets():
    from pathlib import Path
    assets = Path(config.BASE_DIR) / "assets"
    assert Path(config.HOW_TO_VIDEO) == assets / "how_to_download.mp4"
    assert Path(config.HOW_TO_PRICE_VIDEO) == assets / "how_to_price.mp4"
    for p in (config.HOW_TO_VIDEO, config.HOW_TO_PRICE_VIDEO):
        data = Path(p).read_bytes()
        assert b"ftyp" in data[:32] and b"avc1" in data, f"{p}: MP4 с H.264"
        assert b"mp4a" not in data, f"{p}: без звука"
        assert len(data) <= 10 * 1024 * 1024
    assert not (assets / "how_to_price.gif").exists(), "старый GIF заменён на MP4"
    assert not (Path(config.BASE_DIR) / "assetshow_to_price.mp4").exists()


def test_price_caption_exact():
    assert T.ASK_PRICE == (
        "Укажи <b>цену реализации</b> товара, ₽ — сумму, по которой продавец продаёт товар. "
        "С неё считается ДРР.\n"
        "Где смотреть: «Товары и цены» → «Цены и скидки» → колонка <b>«Цена со скидкой»</b> 👆")


WELCOME_EXACT = """👋 Привет!

Это бесплатный бот <a href="https://click-i.ru">Clicki</a> — сервиса аналитики и управления рекламными кампаниями на Wildberries.

Сколько рекламного бюджета ты тратишь на ключи, которые не приносят заказов?

За пару минут бот проверит твою рекламу на WB и покажет:
🔴 какие кластеры сливают деньги (их можно отключить)
🟡 какие стоит докрутить
🟢 какие работают как надо

Нажми «Новый анализ», и начнём 👇"""


async def test_welcome_exact(h):
    r = await h.text("/start")
    assert texts(r) == [WELCOME_EXACT], "приветствие дословно, без добавлений"
    assert T.WELCOME == WELCOME_EXACT
    [msg] = [c for c in r if isinstance(c, SendMessage)]
    assert msg.reply_markup.keyboard[0][0].text == BTN_NEW
    assert "Инструкци" not in msg.text
    assert msg.link_preview_options is not None and msg.link_preview_options.is_disabled, \
        "превью ссылки на сайт отключено"


# ---------- кнопки ДРР ----------

def _drr_rows(markup) -> list[list[str]]:
    return [[b.text for b in row] for row in markup.inline_keyboard]


async def test_drr_buttons_1_to_10(h):
    await h.text(BTN_NEW)
    await h.doc()
    await h.cb("skip_label")
    r = await h.text("9 500")
    [kb] = [c.reply_markup for c in r if isinstance(c, SendMessage) and c.reply_markup]
    assert _drr_rows(kb) == [
        ["1 %", "2 %", "3 %", "4 %", "5 %"],
        ["6 %", "7 %", "8 %", "9 %", "10 %"],
        ["Свой"],
    ]
    assert sum(len(row) for row in kb.inline_keyboard) == 11
    r = await h.cb(buttons(r)["8 %"])
    assert "норма ДРР 8 %" in summary(r)


def test_drr_kb_default_star():
    from bot.keyboards import drr_kb
    rows = _drr_rows(drr_kb(4.0))
    assert rows[0] == ["1 %", "2 %", "3 %", "⭐ 4 %", "5 %"] and len(rows) == 3
    assert sum("⭐" in t for row in rows for t in row) == 1
    assert _drr_rows(drr_kb(10))[1][-1] == "⭐ 10 %"


@pytest.mark.parametrize("default,label,value", [
    (4.5, "⭐ 4,5 %", "drr:4.5"), (4.25, "⭐ 4,25 %", "drr:4.25"), (12, "⭐ 12 %", "drr:12"),
])
def test_drr_kb_fractional_default_on_top(default, label, value):
    from bot.keyboards import drr_kb
    kb = drr_kb(default)
    assert [[(b.text, b.callback_data) for b in row] for row in kb.inline_keyboard][0] == [(label, value)]
    rows = _drr_rows(kb)
    assert rows[1:] == [["1 %", "2 %", "3 %", "4 %", "5 %"], ["6 %", "7 %", "8 %", "9 %", "10 %"],
                        ["Свой"]]


async def test_drr_fractional_default_in_flow(h):
    await h.db.set_default(UID, drr=4.5)
    await h.text(BTN_NEW)
    await h.doc()
    await h.cb("skip_label")
    r = await h.text("9 500")
    assert _drr_rows([c for c in r if isinstance(c, SendMessage)][-1].reply_markup)[0] == ["⭐ 4,5 %"]
    r = await h.cb(buttons(r)["⭐ 4,5 %"])
    assert "норма ДРР 4,5 %" in summary(r)


# ---------- удаление видео-подсказок ----------

def deleted(calls) -> list[int]:
    return [c.message_id for c in calls if isinstance(c, DeleteMessage)]


async def _file_step(h):
    """«➕ Новый анализ» → message_id видео how_to_download."""
    [a] = animations(await h.text(BTN_NEW))
    return h.session.message_id(a)


async def _price_step(h):
    """До шага цены → message_id видео how_to_price."""
    await _file_step(h)
    await h.doc()
    [a] = animations(await h.cb("skip_label"))
    return h.session.message_id(a)


async def test_file_video_deleted_after_file(h):
    vid = await _file_step(h)
    r = await h.doc()
    assert deleted(r) == [vid], "видео шага удалено"
    t = texts(r)
    assert t[0] == f"📄 Файл: {h.fixture_name}", "вместо видео — строка-итог"
    assert "334 кластеров" in t[1] and "Как подписать анализ" in t[2]
    # сначала удаление видео, потом строка-итог
    kinds = [type(c) for c in r if isinstance(c, (DeleteMessage, SendMessage))]
    assert kinds[:2] == [DeleteMessage, SendMessage]


async def test_file_step_bad_input_keeps_video(h):
    await _file_step(h)
    r = await h.doc("report.csv")
    assert deleted(r) == [] and animations(r) == []
    assert texts(r) == ["Это не Excel-файл, пришли выгрузку „Топ поисковых кластеров“"]
    r = await h.text("привет")
    assert deleted(r) == [] and animations(r) == [], "видео не удаляется и не шлётся повторно"
    # правильный файл после ошибок — удаляется то же самое, первое видео
    vid = h.session.message_id(animations(h.session.calls)[0])
    assert deleted(await h.doc()) == [vid]


async def test_price_video_deleted_after_price(h):
    vid = await _price_step(h)
    r = await h.text("дорого")
    assert deleted(r) == [] and animations(r) == [], "ошибка ввода — видео остаётся"
    assert texts(r) == ["Не понял цену. Пришли число, например 9500 или 9 500,50"]
    r = await h.text("0")
    assert deleted(r) == [] and animations(r) == []
    r = await h.text("9 500,5")
    assert deleted(r) == [vid]
    t = texts(r)
    assert t[0] == "💰 Цена реализации: 9 500,5 ₽" and "ДРР" in t[1]


async def test_price_video_deleted_after_default_price_button(h):
    await h.db.set_default(UID, price=9500)
    vid = await _price_step(h)
    r = await h.cb("use_price")
    assert deleted(r) == [vid]
    assert texts(r)[0] == "💰 Цена реализации: 9 500 ₽" and "ДРР" in texts(r)[1]


async def test_full_flow_leaves_no_step_videos(h):
    r_all = []
    r_all += await h.text(BTN_NEW)
    r_all += await h.doc()
    r_all += await h.text("Бомбер")
    r_all += await h.text("9 500")
    r = await h.cb("drr:4")
    sent = {h.session.message_id(a) for a in animations(h.session.calls)}
    assert sent and sent <= set(deleted(h.session.calls)), "оба видео шагов удалены"
    assert "📊 Анализ: Бомбер" in summary(r)
    t = texts(h.session.calls)
    assert f"📄 Файл: {h.fixture_name}" in t and "💰 Цена реализации: 9 500 ₽" in t


@pytest.mark.parametrize("step", ["file", "price"])
async def test_cancel_deletes_step_video(h, step):
    vid = await (_file_step(h) if step == "file" else _price_step(h))
    r = await h.text(BTN_CANCEL)
    assert deleted(r) == [vid]
    assert texts(r) == ["Отменено."], "при отмене строки-итога нет"
    # повторная отмена ничего не удаляет (id сброшен вместе с состоянием)
    assert deleted(await h.text(BTN_CANCEL)) == []


@pytest.mark.parametrize("exit_via", [
    "/start", "/cancel", BTN_HISTORY, BTN_SETTINGS, BTN_HELP, "cb:menu",
])
async def test_exit_to_menu_deletes_step_video(h, exit_via):
    vid = await _price_step(h)
    r = await (h.cb("menu") if exit_via == "cb:menu" else h.text(exit_via))
    assert vid in deleted(r)
    # сценарий сброшен: ввод цены больше не воспринимается как шаг
    assert "💰" not in "\n".join(texts(await h.text("9500")))


async def test_new_analysis_again_deletes_previous_video(h):
    vid = await _file_step(h)
    r = await h.text(BTN_NEW)
    assert deleted(r) == [vid]
    [a] = animations(r)
    assert h.session.message_id(a) != vid


async def test_recalc_price_video_deleted(h):
    await new_analysis(h)
    r = await h.cb("an:recalc:1")
    [a] = animations(r)
    r = await h.text("9 000")
    assert deleted(r) == [h.session.message_id(a)]
    assert texts(r)[0] == "💰 Цена реализации: 9 000 ₽"


async def test_fallback_text_is_deleted_too(h):
    # видео нет — вместо него текст подписи, и он тоже убирается при переходе шага
    h.video.unlink()
    r = await h.text(BTN_NEW)
    [msg] = [c for c in r if isinstance(c, SendMessage)]
    assert deleted(await h.doc()) == [h.session.message_id(msg)]


async def test_delete_error_does_not_break_flow(h, caplog):
    h.session.fail_delete = True
    await _file_step(h)
    with caplog.at_level("DEBUG", logger="bot.cleanup"):
        r = await h.doc()
    assert len(deleted(r)) == 1, "попытка удаления была"
    assert texts(r)[0] == f"📄 Файл: {h.fixture_name}" and "334 кластеров" in texts(r)[1]
    assert "Не удалось удалить сообщение" in caplog.text
    assert all(rec.levelname == "DEBUG" for rec in caplog.records if "удалить" in rec.message)
    r = await h.cb("skip_label")
    r = await h.text("9 500")
    assert texts(r)[0] == "💰 Цена реализации: 9 500 ₽"
    r = await h.cb("drr:4")
    assert "📊 Анализ" in summary(r), "сценарий дошёл до конца"
    assert texts(await h.text(BTN_CANCEL)) == ["Отменено."]


# ---------- удаление прошлой инструкции ----------

def _help_ids(h, r) -> list[int]:
    return [h.session.message_id(c) for c in r if isinstance(c, (SendAnimation, SendMessage))]


async def test_repeat_instruction_deletes_previous(h):
    r1 = await h.text(BTN_HELP)
    ids1 = _help_ids(h, r1)
    assert len(ids1) == 3 and deleted(r1) == []
    assert await h.db.get_help_messages(UID) == (UID, ids1)

    r2 = await h.text("/help")
    assert deleted(r2) == ids1, "сначала удалены оба видео и текст прошлой инструкции"
    first_send = next(i for i, c in enumerate(r2) if isinstance(c, (SendAnimation, SendMessage)))
    assert all(isinstance(c, DeleteMessage) for c in r2[:first_send]), "удаление — до отправки"
    ids2 = _help_ids(h, r2)
    assert len(ids2) == 3 and set(ids2).isdisjoint(ids1)
    assert await h.db.get_help_messages(UID) == (UID, ids2)


async def test_instruction_ids_survive_restart(h):
    ids1 = _help_ids(h, await h.text(BTN_HELP))
    await h.start()  # «перезапуск»: новый Dispatcher, та же БД
    r = await h.text(BTN_HELP)
    assert deleted(r) == ids1


async def test_instruction_delete_error_does_not_break(h, caplog):
    await h.text(BTN_HELP)
    h.session.fail_delete = True
    with caplog.at_level("DEBUG", logger="bot.cleanup"):
        r = await h.text(BTN_HELP)
    assert len(deleted(r)) == 3
    assert texts(r) == [T.HOWTO_CAPTION, T.PRICE_CAPTION, T.INSTRUCTION], "новая инструкция пришла"
    assert caplog.text.count("Не удалось удалить сообщение") == 3


async def test_instruction_during_analysis_deletes_step_video(h):
    vid = await _file_step(h)
    r = await h.text(BTN_HELP)
    assert deleted(r) == [vid]


# ---------- база пользователей, воронка, контакты ----------

from datetime import timedelta  # noqa: E402

from aiogram.types import Contact  # noqa: E402

ADMIN_ID = 1
ADMIN = User(id=ADMIN_ID, is_bot=False, first_name="Админ", username="boss")


async def _as_admin(h, monkeypatch):
    await rebuild(h, monkeypatch, allowed={UID}, admins={ADMIN_ID})


async def contact(h, phone="79991234567", user_id=UID, sender=USER):
    return await h._feed(message=Message(
        message_id=next(h.upd), date=datetime.now(), chat=CHAT, from_user=sender,
        contact=Contact(phone_number=phone, first_name="T", user_id=user_id)))


async def test_events_across_scenario(h):
    await h.text("/start")
    await h.text(BTN_HELP)
    await h.text(BTN_NEW)
    await h.doc("report.csv")
    await h.doc()
    await h.text("Бомбер")
    await h.text("9 500")
    r = await h.cb("drr:4")
    await h.cb("an:recalc:1")
    await h.text(BTN_CANCEL)
    await h.text(BTN_HISTORY)
    await h.cb("an:xl:1")  # «📥 Excel» из истории
    ev = await h.db.events_of(UID)
    names = [e for e, _ in ev]
    assert names == ["start", "instruction_open", "new_analysis", "file_error", "file_ok", "label",
                     "price", "drr", "result_sent", "excel_sent", "recalc", "cancel",
                     "history_open", "excel_download"]
    meta = dict(ev)
    assert meta["file_error"]["reason"] == "не .xlsx"
    assert meta["file_ok"]["clusters"] == 334
    assert meta["label"] == {"label": "Бомбер"} and meta["price"] == {"price": 9500}
    assert meta["drr"] == {"drr": 4.0} and meta["result_sent"]["analysis_id"] == 1
    assert meta["cancel"] == {"step": "price"}, "отмена на шаге цены пересчёта"

    u = await h.db.get_user(UID)
    assert u.last_step == "excel_download" and u.last_step_at is not None
    assert (u.first_name, u.username) == ("T", None)
    assert u.first_seen is not None and u.last_seen >= u.first_seen and u.analyses == 1


async def test_profile_updated_on_each_visit(h):
    await h.text("/start")
    renamed = User(id=UID, is_bot=False, first_name="Тимур", last_name="К", username="timur",
                   language_code="ru")
    await h.text("/start", user=renamed)
    u = await h.db.get_user(UID)
    assert (u.first_name, u.last_name, u.username, u.language_code) == ("Тимур", "К", "timur", "ru")


async def test_track_failure_does_not_break_analysis(h, monkeypatch, caplog):
    async def boom(*a, **kw):
        raise RuntimeError("БД недоступна")
    monkeypatch.setattr(h.db, "add_event", boom)
    with caplog.at_level("WARNING"):
        r = await new_analysis(h)
    assert "📊 Анализ: ИП Иванов" in summary(r), "анализ дошёл до конца"
    assert any(isinstance(c, SendDocument) for c in r)
    assert "не удалось записать событие" in caplog.text


async def test_new_user_notifies_admins(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    r = await h.text("/start")
    assert admin_texts(r, ADMIN_ID) == ["👤 Новый пользователь: T (без username)"]
    assert admin_texts(await h.text("/start"), ADMIN_ID) == [], "только при первом визите"
    monkeypatch.setattr(config, "NOTIFY_NEW_USERS", False)
    other = User(id=77, is_bot=False, first_name="Новый", username="new")
    monkeypatch.setattr(config, "ALLOWED_USER_IDS", set())
    await rebuild(h, monkeypatch)
    assert admin_texts(await h.text("/start", user=other), ADMIN_ID) == [], "флаг выключен"


# ---------- контакт после результата ----------

def _offer(calls) -> SendMessage:
    [m] = [c for c in calls if isinstance(c, SendMessage) and c.text == T.LEAD_ASK]
    return m


async def test_review_offer_after_every_analysis(h):
    r = await new_analysis(h)
    t = texts(r)
    assert t[-1] == T.LEAD_ASK, "после сводки — предложение разбора"
    assert t.index(summary(r)) < len(t) - 1
    m = _offer(r)
    assert [[b.text for b in row] for row in m.reply_markup.inline_keyboard] == \
        [["🙋 Хочу бесплатный разбор"]], "одна inline-кнопка"
    assert "Написать специалисту" not in str(m.reply_markup)
    # и после пересчёта тоже — без ограничения «раз в 7 дней»
    await h.cb("an:recalc:1")
    await h.text("9 000")
    r = await h.cb("drr:4")
    assert texts(r)[-1] == T.LEAD_ASK


async def test_review_click_without_phone_asks_contact(h, monkeypatch):
    monkeypatch.setattr(config, "PRIVACY_URL", "https://click-i.ru/privacy")
    await new_analysis(h)
    r = await h.cb("lead_review")
    [m] = [c for c in r if isinstance(c, SendMessage)]
    assert m.text == (T.LEAD_CONTACT_ASK + "\nНажимая кнопку, ты соглашаешься на обработку "
                                           "персональных данных: https://click-i.ru/privacy")
    kb = m.reply_markup.keyboard
    assert kb[0][0].text == "📱 Оставить контакт" and kb[0][0].request_contact is True
    assert kb[1][0].text == "Не сейчас"
    assert ("lead_request", None) in await h.db.events_of(UID)
    assert (await h.db.get_user(UID)).lead_requested_at is not None


async def test_review_contact_ask_without_privacy_url(h, monkeypatch):
    monkeypatch.setattr(config, "PRIVACY_URL", "")
    await new_analysis(h)
    r = await h.cb("lead_review")
    assert texts(r) == [T.LEAD_CONTACT_ASK]


async def test_lead_contact_saved_and_admins_notified(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    await new_analysis(h)
    await h.cb("lead_review")
    r = await contact(h)
    assert texts(r) == ["Спасибо! Скоро напишем 🙌"]
    [m] = [c for c in r if isinstance(c, SendMessage) and c.chat_id == UID]
    assert m.reply_markup.keyboard[0][0].text == BTN_NEW, "вернули главное меню"
    assert (await h.db.get_user(UID)).phone == "+79991234567"
    assert admin_texts(r, ADMIN_ID) == [
        "🔥 Новый лид: T (без username) +79991234567, анализов: 1, последний: ИП Иванов"]
    assert ("lead_contact", None) in await h.db.events_of(UID)


async def test_review_click_with_phone_is_instant_request(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    await new_analysis(h)
    await contact(h)  # телефон уже сохранён
    # следующий анализ — кнопка снова есть; нажатие сразу создаёт заявку, телефон не спрашиваем
    await h.cb("an:recalc:1")
    await h.text("9 000")
    r = await h.cb("drr:4")
    assert texts(r)[-1] == T.LEAD_ASK
    r = await h.cb("lead_review")
    assert texts(r) == ["Спасибо! Скоро напишем 🙌"]
    assert not any("Оставить контакт" in str(getattr(c, "reply_markup", "")) for c in r)
    assert admin_texts(r, ADMIN_ID) == [
        "🔥 Заявка на разбор: T (без username) +79991234567, анализов: 2, последний: ИП Иванов"]
    ev = [e for e, _ in await h.db.events_of(UID)]
    assert ev[-1] == "lead_request"


async def test_foreign_contact_rejected(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    await new_analysis(h)
    await h.cb("lead_review")
    r = await contact(h, phone="+70000000000", user_id=999)  # переслал чужой контакт
    assert texts(r) == [T.LEAD_FOREIGN_CONTACT]
    assert (await h.db.get_user(UID)).phone is None
    assert admin_texts(r, ADMIN_ID) == []
    r = await contact(h, phone="+70000000000", user_id=None)  # контакт без user_id
    assert texts(r) == [T.LEAD_FOREIGN_CONTACT] and (await h.db.get_user(UID)).phone is None


async def test_lead_later_button(h):
    await new_analysis(h)
    await h.cb("lead_review")
    r = await h.text("Не сейчас")
    assert texts(r) == [T.LEAD_LATER]
    assert (await h.db.get_user(UID)).phone is None


def test_no_manager_username():
    import config as cfg
    assert not hasattr(cfg, "MANAGER_USERNAME")
    assert "MANAGER_USERNAME" not in Path(cfg.BASE_DIR, ".env.example").read_text(encoding="utf-8")


# ---------- /stats, /users, /stuck ----------

async def _seed_funnel(db, now):
    """10 человек на старте, 8 → анализ, 6 → файл, 5 → цена, 4 → ДРР, 4 → результат,
    2 → Excel, 1 → контакт. Плюс старые события (40 дней назад) и ошибки файла."""
    steps = [("start", 10), ("new_analysis", 8), ("file_ok", 6), ("price", 5), ("drr", 4),
             ("result_sent", 4), ("excel_sent", 2), ("lead_request", 2),
             ("lead_contact", 1)]
    for event, n in steps:
        for uid in range(100, 100 + n):
            await db.add_event(uid, event, at=now - timedelta(days=2))
            await db.add_event(uid, event, at=now - timedelta(days=1))  # дубль не считается
    await db.add_event(500, "start", at=now - timedelta(days=40))
    for uid, reason in ((101, "не .xlsx"), (102, "не .xlsx"), (103, "ошибка чтения файла")):
        await db.add_event(uid, "file_error", {"reason": reason}, at=now - timedelta(days=3))
    await db.add_event(104, "materials_open", at=now - timedelta(days=1))
    await db.add_event(600, "start", at=now - timedelta(days=20))


async def test_stats_funnel(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    now = datetime.now()
    await _seed_funnel(h.db, now)
    assert (await h.db.funnel(now - timedelta(days=7)))["start"] == 10
    stats = texts(await h.text("/stats", user=ADMIN), chat_id=UID)[0]
    week, month = stats.split("<b>Воронка за 7 дн.</b>")[1].split("<b>Воронка за 30 дн.</b>")
    assert ("/start: 10\nновый анализ: 8 (80 %)\nфайл принят: 6 (75 %)\nцена: 5 (83 %)\n"
            "ДРР: 4 (80 %)\nполучил результат: 4 (100 %)\nполучил Excel: 2 (50 %)\n"
            "хочет разбор: 2 (100 %)\nоставил контакт: 1 (50 %)") in week
    assert "📚 Открыли материалы: 1" in week
    assert "Топ ошибок файла:\n• не .xlsx — 2\n• ошибка чтения файла — 1" in week
    assert "/start: 11" in month, "за 30 дней + пользователь 20 дней назад, 40 дней — нет"


async def test_admin_commands_hidden_from_users(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    for cmd in ("/stats", "/users", "/stuck"):
        r = await h.text(cmd)
        assert texts(r) == ["Выбери действие в меню."], cmd
        assert not any(isinstance(c, SendDocument) for c in r)


async def test_users_excel(h, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    await _as_admin(h, monkeypatch)
    await new_analysis(h)
    await contact(h)
    await h.db.upsert_user(55, "stuck_guy", "Застрявший", None, "ru")
    await h.db.add_event(55, "price", at=datetime.now() - timedelta(hours=2))
    r = await h.text("/users", user=ADMIN)
    [d] = [c for c in r if isinstance(c, SendDocument)]
    assert d.document.filename.startswith("users_") and "Пользователей: 3" in d.caption
    ws = load_workbook(BytesIO(d.document.data)).active
    assert [c.value for c in ws[1]] == ["ID", "Username", "Имя", "Телефон", "Первый визит",
                                        "Последний визит", "Анализов", "Последний шаг", "Когда",
                                        "Застрял"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    me, stuck = rows[UID], rows[55]
    assert me[3].value == "+79991234567" and me[6].value == 1 and me[9].value == "нет"
    assert me[7].value == "оставил контакт"
    assert stuck[1].value == "@stuck_guy" and stuck[1].hyperlink.target == "https://t.me/stuck_guy"
    assert stuck[7].value == "цена" and stuck[9].value == "да"


async def test_stuck_list(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    now = datetime.now()
    await h.db.upsert_user(55, "stuck_guy", "Застрявший", None, "ru")
    await h.db.add_event(55, "price", at=now - timedelta(hours=2))          # застрял
    await h.db.upsert_user(56, None, "Свежий", None, "ru")
    await h.db.add_event(56, "file_ok", at=now - timedelta(minutes=5))      # ещё думает
    await h.db.upsert_user(57, "done", "Дошёл", None, "ru")
    await h.db.add_event(57, "result_sent", at=now - timedelta(hours=3))    # дошёл
    await h.db.upsert_user(58, "old", "Давно", None, "ru")
    await h.db.add_event(58, "price", at=now - timedelta(days=10))          # старше 7 дней
    await h.db.upsert_user(59, "mat", "Материалы", None, "ru")
    await h.db.add_event(59, "materials_open", at=now - timedelta(hours=3))  # после результата
    [t] = texts(await h.text("/stuck", user=ADMIN), chat_id=UID)
    assert "Застряли за 7 дн.:</b> 1" in t
    assert "• Застрявший @stuck_guy — цена," in t
    for name in ("Свежий", "Дошёл", "Давно", "Материалы"):
        assert name not in t


# ---------- «📚 Полезные материалы» ----------

def documents(calls) -> list[SendDocument]:
    return [c for c in calls if isinstance(c, SendDocument) and c.caption]


def _rows(markup) -> list[list[str]]:
    return [[b.text for b in row] for row in markup.inline_keyboard]


def _reply_rows(markup) -> list[list[str]]:
    return [[b.text for b in row] for row in markup.keyboard]


async def test_main_menu_layout(h):
    [m] = [c for c in await h.text("/start") if isinstance(c, SendMessage)]
    assert _reply_rows(m.reply_markup) == [
        ["➕ Новый анализ"],
        ["📂 История", "⚙️ Настройки"],
        ["📖 Инструкция", "📚 Полезные материалы"],
    ]


async def test_materials_not_under_summary_nor_history(h):
    r = await new_analysis(h)
    [m] = [c for c in r if isinstance(c, SendMessage) and c.text.startswith("<b>📊 Анализ")]
    assert _rows(m.reply_markup) == [["🔁 Пересчитать", "🏠 Меню"]]
    r = await h.text(BTN_HISTORY)
    r = await h.cb(buttons(r)["1"])
    [m] = [c for c in r if isinstance(c, SendMessage) and c.reply_markup]
    rows = _rows(m.reply_markup)
    assert rows[0] == ["📥 Excel", "🔁 Пересчитать"] and "🗑 Удалить" in rows[1]
    assert "📚 Полезные материалы" not in sum(rows, [])


async def test_materials_from_main_menu(h):
    r = await h.text("📚 Полезные материалы")
    assert texts(r) == [T.MATERIALS_INTRO]
    assert [d.caption for d in documents(r)] == [m["title"] for m in T.MATERIALS]
    assert all(isinstance(d.document, FSInputFile) for d in documents(r))
    assert ("materials_open", None) in await h.db.events_of(UID)
    # повтор в течение минуты — файлы не шлём, отвечаем текстом
    h.clock[0] += 30
    r = await h.text("📚 Полезные материалы")
    assert documents(r) == [] and texts(r) == ["Материалы уже выше 👆"]
    # через минуту — снова, уже по file_id
    h.clock[0] += 31
    docs = documents(await h.text("📚 Полезные материалы"))
    assert [d.document for d in docs] == ["Разбор_Поиск#1", "Разбор_Полки#1"]


async def test_materials_menu_button_leaves_flow(h):
    # кнопка меню во время анализа: видео шага убирается, сценарий сбрасывается
    [a] = animations(await h.text(BTN_NEW))
    vid = h.session.message_id(a)
    r = await h.text("📚 Полезные материалы")
    assert vid in [c.message_id for c in r if isinstance(c, DeleteMessage)]
    assert len(documents(r)) == 2
    assert "Файл принят" not in "\n".join(texts(await h.doc())), "шаг «пришли файл» сброшен"


async def test_materials_sent_then_by_file_id(h):
    await new_analysis(h)
    r = await h.cb("materials")
    intro = [c for c in r if isinstance(c, SendMessage)]
    assert [c.text for c in intro] == [T.MATERIALS_INTRO]
    assert intro[0].text.startswith('📚 Полезные материалы от <a href="https://click-i.ru">Clicki</a> — ')
    assert intro[0].link_preview_options.is_disabled, "превью ссылки отключено"
    docs = documents(r)
    assert [d.caption for d in docs] == [m["title"] for m in T.MATERIALS]
    assert all(isinstance(d.document, FSInputFile) for d in docs), "первый раз — загрузка"
    assert [Path(d.document.path).name for d in docs] == ["Разбор_Поиск.pdf", "Разбор_Полки.pdf"]
    assert ("materials_open", None) in await h.db.events_of(UID)

    h.clock[0] += 61  # прошла минута — можно снова
    docs = documents(await h.cb("materials"))
    assert [d.document for d in docs] == ["Разбор_Поиск#1", "Разбор_Полки#1"], "по file_id"
    assert h.session.uploads["Разбор_Поиск"] == 1 and h.session.uploads["Разбор_Полки"] == 1


async def test_materials_repeat_within_minute(h):
    await h.cb("materials")
    h.clock[0] += 30
    r = await h.cb("materials")
    assert documents(r) == [] and not [c for c in r if isinstance(c, SendMessage)]
    [ans] = [c for c in r if isinstance(c, AnswerCallbackQuery)]
    assert ans.text == "Материалы уже выше 👆"
    h.clock[0] += 31  # 61 секунда с первой отправки
    assert len(documents(await h.cb("materials"))) == 2


async def test_materials_missing_file_skipped(h, caplog):
    (h.materials_dir / "Разбор_Полки.pdf").unlink()
    with caplog.at_level("WARNING"):
        r = await h.cb("materials")
    assert [d.caption for d in documents(r)] == [T.MATERIALS[0]["title"]]
    assert texts(r) == [T.MATERIALS_INTRO]
    assert "Нет файла материала" in caplog.text


async def test_materials_none_available(h):
    for m in T.MATERIALS:
        (h.materials_dir / m["file"]).unlink()
    r = await h.cb("materials")
    assert texts(r) == ["Материалы скоро появятся"] and documents(r) == []
    # ничего не отправили — повтор не блокируется
    assert texts(await h.cb("materials")) == ["Материалы скоро появятся"]


async def test_materials_changed_file_reuploaded(h):
    await h.cb("materials")
    (h.materials_dir / "Разбор_Поиск.pdf").write_bytes(b"%PDF-1.4 new version, other size")
    h.clock[0] += 61
    docs = documents(await h.cb("materials"))
    assert isinstance(docs[0].document, FSInputFile), "файл изменился — загрузка заново"
    assert docs[1].document == "Разбор_Полки#1"


async def test_materials_stale_file_id_reuploaded(h):
    await h.cb("materials")
    h.session.bad_file_ids.add("Разбор_Поиск#1")
    h.clock[0] += 61
    r = await h.cb("materials")
    poisk = [d for d in documents(r) if d.caption == T.MATERIALS[0]["title"]]
    assert [type(d.document) for d in poisk] == [str, FSInputFile]


async def test_stats_shows_materials(h, monkeypatch):
    await _as_admin(h, monkeypatch)
    await h.cb("materials")
    stats = texts(await h.text("/stats", user=ADMIN))[0]
    assert stats.count("📚 Открыли материалы: 1") == 2, "за 7 и за 30 дней"


def test_real_materials_on_disk():
    from pathlib import Path as P
    for m in T.MATERIALS:
        p = P(config.MATERIALS_DIR) / m["file"]
        assert p.is_file(), p
        assert p.read_bytes()[:5] == b"%PDF-"
        assert len(m["title"]) <= 1024


async def test_excel_auto_sent_and_history_button(h):
    r = await new_analysis(h)
    docs = [c for c in r if isinstance(c, SendDocument)]
    assert [d.document.filename for d in docs] == ["ИП Иванов.xlsx"], "Excel пришёл без нажатия"
    ev = [e for e, _ in await h.db.events_of(UID)]
    assert ev[-2:] == ["result_sent", "excel_sent"]
    # в истории Excel сам не приходит — там кнопка есть и работает
    r = await h.text(BTN_HISTORY)
    r = await h.cb(buttons(r)["1"])
    assert "📥 Excel" in buttons(r) and not [c for c in r if isinstance(c, SendDocument)]
    r = await h.cb(buttons(r)["📥 Excel"])
    assert [c.document.filename for c in r if isinstance(c, SendDocument)] == ["ИП Иванов.xlsx"]
    assert (await h.db.events_of(UID))[-1][0] == "excel_download"


def test_materials_titles():
    assert [m["title"] for m in T.MATERIALS] == [
        "🔍 Почему первое место в поиске не всегда приносит больше продаж",
        "🧩 Как устроены рекомендательные полки WB и почему их используют неэффективно",
    ]


def test_brand_name_in_all_user_texts():
    """Во всех видимых пользователю текстах — «Clicki»; «Click-i.ru» / «click-i» только в URL."""
    import re
    from bot import keyboards, materials, texts as tx
    from core import report
    sources = [tx, keyboards, materials, report,
               __import__("bot.handlers.start", fromlist=["x"]),
               __import__("bot.handlers.new_analysis", fromlist=["x"]),
               __import__("bot.handlers.history", fromlist=["x"]),
               __import__("bot.handlers.settings", fromlist=["x"]),
               __import__("bot.lead", fromlist=["x"])]
    for mod in sources:
        src = Path(mod.__file__).read_text(encoding="utf-8")
        without_urls = re.sub(r"https?://click-i\.ru\S*", "", src)
        assert not re.search(r"(?i)click-i", without_urls), mod.__name__
    assert '<a href="https://click-i.ru">Clicki</a>' in tx.WELCOME
    assert tx.MATERIALS_INTRO.startswith('📚 Полезные материалы от <a href="https://click-i.ru">Clicki</a>')
    assert "Clicki" in tx.LEAD_ASK
